from __future__ import annotations

import json
import logging
import threading
import webbrowser
from datetime import datetime, timezone
from tkinter import ttk, messagebox, filedialog, simpledialog
import tkinter as tk
import customtkinter as ctk
from typing import Any, Optional

from gui.styles import COLORS, FONTS, AutohideScrollbar
from gui.dialogs.edit_form_dialog import EditFormDialog
from gui.dialogs.prompt_editor_dialog import PromptEditorDialog
from gui.panels.diff_panel import DiffPanel
from gui.panels.details_panel import DetailsPanel

logger = logging.getLogger(__name__)

# ─────────────────────────────────────────────────────── table defs ──────────────────
# Each entry: label, icon, section, table (actual SQL table), sql (SELECT),
#   cols [(key, header, width, filterable)], editable, pk,
#   fk_join (column to filter by parent ID, optional),
#   form [(field, label, type)]  where type is text|number|bool|multiline|select:a,b|client_select

TABLES: dict[str, dict[str, Any]] = {
    "all_products": {
        "label": "All Products", "icon": "📦", "section": "raw",
        "table": "all_products",
        "sql": """
            SELECT id, mp, sku, merchant_name AS merchant,
                   CASE avail_code WHEN 1 THEN 'Yes' WHEN 0 THEN 'No' ELSE 'Ltd' END AS avail,
                   category, image,
                   name, url, url_tag,
                   run_at as seen, task_name
            FROM all_products
            ORDER BY run_at DESC LIMIT 10000
        """,
        "cols": [
            ("id","ID",40,False), ("mp","MP",65,True),
            ("sku","SKU",70,True), ("category","Category",120,True),
            ("merchant","Merchant",100,True), 
            ("avail","Avail",50,True),
            ("name","Title",280,True), 
            ("task_name","Task",150,True),
            ("url_tag","Tag",90,True),
            ("url","URL",180,False),
            ("seen","Date",110,False),
        ],
        "editable": False, "pk": "id",
    },
    "raw_sessions": {
        "label": "Scrape Log", "icon": "🔄", "section": "raw",
        "table": "scrape_log",
        "sql": """
            SELECT snapshot_id, run_at as date, task, task_type, product_count as count, status, client
            FROM scrape_log ORDER BY run_at DESC
        """,
        "cols": [
            ("snapshot_id","ID",50,False), ("date","Date",120,True),
            ("task","Task",200,True), ("task_type","Type",80,True),
            ("count","Found",55,False), ("status","Status",85,True),
            ("client","Client",100,True),
        ],
        "editable": False, "pk": "snapshot_id",
    },
    "tasks": {
        "label": "Tasks", "icon": "📋", "section": "business",
        "table": "tasks", "fk_join": "t.client_id",
        "sql": """
            SELECT t.id, t.title, t.task_type, t.schedule_type,
                   (SELECT COUNT(*) FROM snapshots s WHERE s.task_id = t.id) AS snapshot_count,
                   (SELECT MAX(run_at) FROM snapshots s WHERE s.task_id = t.id) AS last_run,
                   DATE(t.created_at) AS created
            FROM tasks t
            ORDER BY t.created_at DESC
        """,
        "cols": [
            ("id", "ID", 40, False), ("title", "Title", 220, True),
            ("task_type", "Type", 100, True),
            ("schedule_type", "Schedule", 100, True),
            ("snapshot_count", "Snapshots", 80, False),
            ("last_run", "Last Run", 150, False),
            ("created", "Created", 85, False),
        ],
        "editable": True, "pk": "id",
        "form": [
            ("client_id", "Client", "client_select"),
            ("title", "Task Title *", "text"),
            ("task_type", "Task Type (tracking|discovery)", "select:tracking,discovery"),
            ("schedule_type", "Schedule", "select:on_demand,daily,weekly"),
            ("query_params", "Scraping Query / Params", "multiline"),
            ("status", "Status", "select:active,paused"),
            ("notes", "Notes", "multiline")
        ]
    },
    "snapshots": {
        "label": "Snapshots", "icon": "📷", "section": "business",
        "table": "snapshots", "fk_join": "s.task_id",
        "sql": """
            SELECT s.id, s.run_at, s.product_count, s.status, t.query_params as scope
            FROM snapshots s
            LEFT JOIN tasks t ON s.task_id = t.id
            ORDER BY s.run_at DESC
        """,
        "cols": [
            ("id", "ID", 40, False), ("run_at", "Run At", 150, False),
            ("product_count", "Products", 80, False), ("status", "Status", 80, True),
            ("scope", "Scope (Params)", 300, True),
        ],
        "editable": False, "pk": "id",
    },
    "snapshot_products": {
        "label": "Snapshot Products", "icon": "📦", "section": "business",
        "table": "snapshot_products", "fk_join": "sp.snapshot_id",
        "sql": """
            SELECT sp.id, 
                   sp.product_id,
                   sp.mp,
                   sp.sku,
                   sp.name, 
                   sp.category,
                   sp.image,
                   sp.price,
                   CASE sp.avail_code WHEN 1 THEN 'Yes' WHEN 0 THEN 'No' ELSE 'Ltd' END AS avail,
                   sp.merchant_name AS merchant, 
                   sp.url_tag,
                   sp.url
            FROM snapshot_products sp
        """,
        "cols": [
            ("id", "ID", 40, False), ("product_id", "P_ID", 80, True),
            ("mp", "MP", 65, True),
            ("sku", "SKU", 70, True), ("name", "Name", 280, True),
            ("category", "Category", 150, True),
            ("price", "Price", 75, True), ("avail", "Avail", 50, True),
            ("merchant", "Merchant", 100, True), ("url_tag", "Tag", 90, True), ("url", "URL", 180, False)
        ],
        "editable": False, "pk": "id",
    },
    "clients": {
        "label": "Clients", "table": "clients", "section": "project",
        "pk": "id",
        "form": [
            ("name", "Client Name *", "text"),
            ("contact_info", "Contact Info / Telegram", "text"),
            ("notes", "Notes", "multiline")
        ]
    }
}

# ------------------------------------------------------- Control Panel -------

class DbBrowserWindow(ctk.CTkToplevel):

    def __init__(self, master, db, scheduler=None):
        super().__init__(master)
        self.db = db
        self.scheduler = scheduler

        self.title("Database Control Panel")
        self.geometry("1440x900")
        self.minsize(1100, 650)

        # state
        self._current_table: str = "all_products"
        self._active_nav_id: Optional[int] = None
        self._all_rows: list[dict] = []
        self._displayed_rows: list[dict] = []
        self._sort_col: Optional[str] = None
        self._sort_dir: int = 1   # 1=ASC -1=DESC
        self._filter_vars: dict[str, tk.StringVar] = {}
        self._global_var = tk.StringVar()
        self._global_var.trace_add("write", self._on_filter_change)
        self._col_keys: list[str] = []
        self._active_nav: Optional[ctk.CTkButton] = None

        self._selected_client_id: Optional[int] = None
        self._selected_task_id: Optional[int] = None
        self._selected_snapshot_id: Optional[int] = None
        
        self._setup_styles()
        self._setup_ui()
        self._refresh_client_sidebar()
        self._load_table("all_products")

        self.grab_set()
        self.after(20, self.lift)
        self.focus_force()

    # ──────────────── styles ──────────────────────────────────────────────────────────

    def _setup_styles(self):
        style = ttk.Style()
        style.theme_use("default")
        bg, fg, sel = "#2d2d2d", "#cccccc", "#1f538d"
        hdr = "#1e1e1e"
        style.configure("DB.Treeview",
            background=bg, foreground=fg, fieldbackground=bg,
            rowheight=24, borderwidth=0, font=("Segoe UI", 11))
        style.configure("DB.Treeview.Heading",
            background=hdr, foreground="#8ab4f8",
            font=("Segoe UI", 11, "bold"), relief="flat", padding=(4, 4))
        style.map("DB.Treeview",
            background=[("selected", sel)],
            foreground=[("selected", "#ffffff")])
        style.map("DB.Treeview.Heading",
            background=[("active", "#333333")])

    # ──────────────── UI layout ───────────────────────────────────────────────────────

    def _setup_ui(self):
        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(0, weight=1)

        # sidebar
        self._sidebar = ctk.CTkScrollableFrame(self, width=260, corner_radius=0,
                                               fg_color=COLORS["sidebar"])
        self._sidebar.grid(row=0, column=0, sticky="nsew")

        # main area
        main = ctk.CTkFrame(self, corner_radius=0, fg_color="transparent")
        main.grid(row=0, column=1, sticky="nsew", padx=0, pady=0)
        main.grid_rowconfigure(2, weight=1)
        main.grid_columnconfigure(0, weight=1)

        self._build_toolbar(main)
        self._build_filter_bar(main)
        self._build_treeview(main)
        self._details_panel = DetailsPanel(main)
        self._build_statusbar(main)

    # ──────────────── sidebar ─────────────────────────────────────────────────────────

    def _sidebar_section(self, text: str):
        ctk.CTkLabel(self._sidebar, text=text, font=("Segoe UI", 10, "bold"),
                     text_color="#888888", anchor="w").pack(fill="x", padx=10, pady=(12, 2))

    def _sidebar_btn(self, text: str, command, indent=0) -> ctk.CTkButton:
        btn = ctk.CTkButton(
            self._sidebar, text=text, anchor="w",
            fg_color="transparent", hover_color="#3a3a3a",
            text_color=COLORS["fg"], font=("Segoe UI", 12),
            height=30, command=command,
        )
        btn.pack(fill="x", padx=(8 + indent, 8), pady=1)
        return btn

    def _set_active_nav(self, btn: ctk.CTkButton):
        if self._active_nav:
            try:
                self._active_nav.configure(fg_color="transparent")
            except Exception:
                pass
        self._active_nav = btn
        btn.configure(fg_color="#1f538d", text_color="#ffffff")

    def _refresh_client_sidebar(self):
        for w in self._sidebar.winfo_children():
            w.destroy()

        # ── Buttons section (Unified Column) ──
        btn_container = ctk.CTkFrame(self._sidebar, fg_color="transparent")
        btn_container.pack(fill="x", padx=8, pady=(8, 4))
        
        def s_btn(text, cmd, color=None):
            b = ctk.CTkButton(
                btn_container, text=text, height=32,
                fg_color=color if color else COLORS["sidebar"],
                font=FONTS["normal"], command=cmd
            )
            b.pack(fill="x", pady=2)
            return b

        s_btn("+ New Client", self._on_new_client, COLORS["success"])
        s_btn("✎ Rename Task", self._on_rename_selected, "#444444")
        s_btn("🗑 Delete", self._on_delete_current, COLORS["error"])

        # ── Clients section ──
        self._sidebar_section("CLIENTS")

        conn = self.db.get_connection()
        try:
            clients = conn.execute(
                "SELECT id, name FROM clients ORDER BY created_at DESC"
            ).fetchall()
        except Exception:
            clients = []
            
        self._client_btns: dict[int, ctk.CTkButton] = {}
        self._task_btns: dict[int, ctk.CTkButton] = {}
        self._raw_btns: dict[str, ctk.CTkButton] = {}
        for c in clients:
            cid = c["id"]
            cname = c["name"]
            
            # Client frame
            c_frame = ctk.CTkFrame(self._sidebar, fg_color="transparent")
            c_frame.pack(fill="x", padx=8, pady=1)

            btn = ctk.CTkButton(
                c_frame, text=f"📂 {cname}", anchor="w",
                fg_color="transparent", hover_color="#3a3a3a",
                text_color=COLORS["accent"], font=FONTS["normal"],
                command=lambda c_id=cid: self._on_nav_client(c_id)
            )
            btn.bind("<Double-1>", lambda e, c_id=cid: self._on_sidebar_double_click(c_id))
            btn.pack(side="left", fill="x", expand=True)
            self._client_btns[cid] = btn

            # Fetch tasks for this client (with snapshot counts)
            try:
                tasks = conn.execute("""
                    SELECT t.*, (SELECT COUNT(id) FROM snapshots s WHERE s.task_id = t.id) as snap_count
                    FROM tasks t WHERE t.client_id = ? ORDER BY t.created_at DESC
                """, (cid,)).fetchall()
            except Exception:
                tasks = []
                
            for t in tasks:
                tid = t["id"]
                ttitle = t["title"]
                
                # Build mask: [RAPE]
                params = json.loads(t["query_params"] or "{}")
                mps = params.get("marketplaces", [])
                mask = ""
                for site_id, char in [("rozetka", "R"), ("allo", "A"), ("prom", "P"), ("epicentrk", "E"), ("hotline", "H")]:
                    if site_id in mps: mask += char
                
                date_str = (t["created_at"] or "")[:10]
                count_str = f"[{t['snap_count']}]"
                mask_str = f"[{mask}]" if mask else ""
                t_type = t["task_type"] if "task_type" in t.keys() else "discovery"
                
                display_label = f"{date_str}  {count_str} {mask_str} {ttitle} [{t_type}]"

                t_row_frame = ctk.CTkFrame(self._sidebar, fg_color="transparent")
                t_row_frame.pack(fill="x", padx=(12, 8), pady=1)
                
                run_btn = ctk.CTkButton(
                    t_row_frame, text="RUN", width=38, height=22,
                    fg_color="#444444", hover_color="#555555",
                    text_color="#ffffff", font=("Segoe UI", 9, "bold"),
                    command=lambda t_id=tid: self._on_rerun_task(t_id)
                )
                run_btn.pack(side="left", padx=(0, 5))

                t_btn = ctk.CTkButton(
                    t_row_frame, text=display_label, anchor="w",
                    fg_color="transparent", hover_color="#3a3a3a",
                    text_color=COLORS["fg"], font=("Segoe UI", 10),
                    height=24, command=lambda t_id=tid: self._on_nav_task(t_id)
                )
                t_btn.bind("<Double-1>", lambda e, t_data=t: self._on_edit_task_modal(t_data))
                t_btn.pack(side="left", fill="x", expand=True)
                self._task_btns[tid] = t_btn


        # --- Raw Data section ---
        self._sidebar_section("RAW DATA")
        self._raw_btns = {}
        for key in ("all_products", "raw_sessions"):
            d = TABLES[key]
            btn = self._sidebar_btn(
                f"{d['icon']}  {d['label']}",
                lambda k=key: self._on_nav_table(k),
            )
            self._raw_btns[key] = btn

    # --- toolbar -----------------------------------------------------------------

    def _build_toolbar(self, parent):
        tb = ctk.CTkFrame(parent, height=46, fg_color="#252526", corner_radius=0)
        tb.grid(row=0, column=0, sticky="ew")
        tb.grid_propagate(False)

        def btn(text, cmd, color=None):
            kw = {"fg_color": color} if color else {}
            ctk.CTkButton(tb, text=text, width=0, height=30,
                          font=("Segoe UI", 11), command=cmd, **kw).pack(
                side="left", padx=4, pady=8)

        btn("➕ Add", self._on_add_row)
        btn("📝 Edit", self._on_edit_row)
        btn("🗑 Delete", self._on_delete_selected, COLORS["error"])
        btn("🔥 Clear", self._on_clear_table, "#8B4513")
        btn("🚀 Rerun", self._on_rerun_selected, "#555555")
        btn("📊 Diff", self._on_compare_snapshots, "#1f538d")
        btn("🤖 AI", self._on_normalize)
        btn("📥 Excel", self._on_export, COLORS["success"])
        btn("🔄 Refresh", self._on_refresh)

        # Export Report button — enabled only when Snapshots table is active
        self._export_report_btn = ctk.CTkButton(
            tb, text="📋 Export Report", width=0, height=30,
            font=("Segoe UI", 11),
            fg_color="#4a3f8a", hover_color="#6a5faa",
            state="disabled",
            command=self._on_export_report,
        )
        self._export_report_btn.pack(side="left", padx=4, pady=8)

        self._stats_lbl = ctk.CTkLabel(tb, text="", font=("Segoe UI", 10),
                                       text_color="#888888")
        self._stats_lbl.pack(side="right", padx=12)

    # ──────────────── filter bar ──────────────────────────────────────────────────────

    def _build_filter_bar(self, parent):
        fb = ctk.CTkFrame(parent, fg_color="#1e1e1e", height=38, corner_radius=0)
        fb.grid(row=1, column=0, sticky="ew")
        fb.grid_propagate(False)

        ctk.CTkLabel(fb, text="Search:", font=("Segoe UI", 11, "bold")).pack(side="left", padx=(8, 2), pady=4)
        self._search_entry = ctk.CTkEntry(
            fb, textvariable=self._global_var,
            placeholder_text="Search all columns...",
            width=260, height=28, font=("Segoe UI", 11))
        self._search_entry.pack(side="left", padx=(0, 16), pady=5)

        # Per-column filters frame (populated dynamically)
        self._col_filter_frame = ctk.CTkFrame(fb, fg_color="transparent")
        self._col_filter_frame.pack(side="left", fill="x", expand=True)

    def _rebuild_col_filters(self, cols):
        for w in self._col_filter_frame.winfo_children():
            w.destroy()
        self._filter_vars = {}
        for key, header, width, filterable in cols:
            if not filterable or width == 0:
                continue
            var = tk.StringVar()
            var.trace_add("write", self._on_filter_change)
            self._filter_vars[key] = var
            cell = ctk.CTkFrame(self._col_filter_frame, fg_color="transparent")
            cell.pack(side="left", padx=3)
            ctk.CTkLabel(cell, text=header, font=("Segoe UI", 9),
                         text_color="#666666").pack(anchor="w")
            ctk.CTkEntry(cell, textvariable=var, width=max(40, min(width, 120)),
                         height=22, font=("Segoe UI", 10)).pack()

    # ──────────────── treeview ────────────────────────────────────────────────────────

    def _build_treeview(self, parent):
        tf = tk.Frame(parent, background="#2d2d2d")
        tf.grid(row=2, column=0, sticky="nsew")
        parent.grid_rowconfigure(2, weight=1)
        parent.grid_columnconfigure(0, weight=1)

        self._tree = ttk.Treeview(tf, style="DB.Treeview",
                                  selectmode="extended", show="headings")
        vsb = AutohideScrollbar(tf, orientation="vertical", command=self._tree.yview)
        hsb = AutohideScrollbar(tf, orientation="horizontal", command=self._tree.xview)
        self._tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        self._tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        tf.grid_rowconfigure(0, weight=1)
        tf.grid_columnconfigure(0, weight=1)

        # Tag styling
        self._tree.tag_configure("irrelevant", foreground="#555555")

        self._tree.bind("<Double-1>", self._on_double_click)
        self._tree.bind("<Delete>", lambda e: self._on_delete_selected())
        self._tree.bind("<Button-3>", self._show_context_menu)
        self._tree.bind("<<TreeviewSelect>>", self._on_tree_select)

    # ──────────────── status bar ──────────────────────────────────────────────────────

    def _build_statusbar(self, parent):
        sb = ctk.CTkFrame(parent, height=26, fg_color="#1a1a1a", corner_radius=0)
        sb.grid(row=4, column=0, sticky="ew")
        sb.grid_propagate(False)
        self._status_lbl = ctk.CTkLabel(sb, text="Ready", font=("Segoe UI", 10),
                                        text_color="#666666", anchor="w")
        self._status_lbl.pack(side="left", padx=10)

    def _set_status(self, msg: str):
        try:
            self._status_lbl.configure(text=msg)
        except Exception:
            pass

    # ──────────────── navigation ──────────────────────────────────────────────────────

    def _on_nav_table(self, key: str):
        self._selected_client_id = None
        self._selected_task_id = None
        self._selected_snapshot_id = None
        if hasattr(self, "_raw_btns") and key in self._raw_btns:
            self._set_active_nav(self._raw_btns[key])
        self._load_table(key)

    def _on_nav_client(self, client_id: int):
        self._selected_client_id = client_id
        if client_id in self._client_btns:
            self._set_active_nav(self._client_btns[client_id])
        self._selected_task_id = None
        self._selected_snapshot_id = None
        self._load_table("tasks")

    def _on_nav_task(self, task_id: int):
        self._selected_task_id = task_id
        self._selected_client_id = None
        self._selected_snapshot_id = None
        if hasattr(self, "_task_btns") and task_id in self._task_btns:
            self._set_active_nav(self._task_btns[task_id])
        self._load_table("snapshots")

    def _on_nav_snapshot(self, snapshot_id: int):
        self._selected_snapshot_id = snapshot_id
        self._load_table("snapshot_products")

    def _on_delete_client(self, client_id: int):
        if not messagebox.askyesno("Delete Client", "Are you sure? All related tasks and snapshots will be permanently deleted."):
            return
        try:
            conn = self.db.get_connection()
            conn.execute("DELETE FROM clients WHERE id = ?", (client_id,))
            conn.commit()
            if self._selected_client_id == client_id:
                self._on_nav_table("all_products")
            self._refresh_client_sidebar()
        except Exception as e:
            messagebox.showerror("Error", str(e))

    def _on_delete_task(self, task_id: int):
        if not messagebox.askyesno("Delete Task", "Are you sure? This will delete the task and its snapshots."):
            return
        try:
            conn = self.db.get_connection()
            conn.execute("DELETE FROM tasks WHERE id = ?", (task_id,))
            conn.commit()
            if self._selected_task_id == task_id:
                self._on_nav_table("all_products")
            self._refresh_client_sidebar()
        except Exception as e:
            messagebox.showerror("Error", str(e))

    def _on_sidebar_double_click(self, client_id: int):
        self._current_table = "clients"
        try:
            conn = self.db.get_connection()
            row = conn.execute("SELECT * FROM clients WHERE id = ?", (client_id,)).fetchone()
            if row:
                self._show_edit_form(row_data=dict(row))
        except Exception as e:
            messagebox.showerror("Error", str(e))

    def _on_rerun_task(self, task_id: int):
        """Extract params from task and trigger MainWindow to re-run it."""
        try:
            conn = self.db.get_connection()
            row = conn.execute("SELECT client_id, query_params FROM tasks WHERE id = ?", (task_id,)).fetchone()
            if not row: return
            
            c_id = row["client_id"]
            params = json.loads(row["query_params"])
            
            if hasattr(self.master, "active_client_id"):
                self.master.active_client_id = c_id
            
            discovery_tasks = []
            # Extract queries and marketplaces
            for q in params.get("queries", []):
                for mp in params.get("marketplaces", []):
                    discovery_tasks.append((mp, "MAPI", q, None))
            
            pages = params.get("pages", 1)
            threads = params.get("threads", 1)
            delay = params.get("delay", 1.5)
            skip_stock = params.get("skip_stock", True)
            
            if hasattr(self.master, "_run_discovery_batch"):
                # Call with task_id to reuse the task record
                self.master._run_discovery_batch(
                    discovery_tasks, pages, skip_stock, 
                    threads=threads, delay=delay, task_id=task_id
                )
                messagebox.showinfo("Started", f"Re-running task {task_id}. New snapshot will be added.")
            else:
                messagebox.showerror("Error", "Main Window not found.")
        except Exception as e:
            messagebox.showerror("Error", str(e))

    def _on_delete_current(self):
        """Intelligently delete whatever is currently selected/browsed."""
        if self._selected_snapshot_id:
            msg = f"Delete snapshot #{self._selected_snapshot_id}?"
            target = ("snapshots", "id", self._selected_snapshot_id)
        elif self._selected_task_id:
            msg = f"Delete active task #{self._selected_task_id} and all its snapshots?"
            target = ("tasks", "id", self._selected_task_id)
        elif self._selected_client_id:
            msg = f"Delete client #{self._selected_client_id} and ALL their tasks/snapshots?"
            target = ("clients", "id", self._selected_client_id)
        else:
            messagebox.showinfo("Info", "Select a client or task in the sidebar first.")
            return

        if not messagebox.askyesno("Confirm Delete", msg): return
        
        try:
            conn = self.db.get_connection()
            conn.execute(f"DELETE FROM {target[0]} WHERE {target[1]} = ?", (target[2],))
            conn.commit()
            
            # Reset state
            if target[0] == "clients": self._selected_client_id = None
            if target[0] in ("clients", "tasks"): self._selected_task_id = None
            self._selected_snapshot_id = None
            
            self._refresh_client_sidebar()
            self._on_nav_table("all_products")
            self._set_status("Item deleted.")
        except Exception as e:
            messagebox.showerror("Error", str(e))

    def _on_rename_selected(self):
        if not self._selected_task_id:
            messagebox.showinfo("Info", "Select a task in the sidebar first.")
            return
        
        new_name = simpledialog.askstring("Rename Task", "Enter new custom name for this task:",
                                          initialvalue="")
        if new_name is None: return # Cancelled
        
        try:
            conn = self.db.get_connection()
            conn.execute("UPDATE tasks SET title = ? WHERE id = ?", (new_name, self._selected_task_id))
            conn.commit()
            self._refresh_client_sidebar()
            self._set_status("Task renamed.")
        except Exception as e:
            messagebox.showerror("Error", str(e))

    def _on_edit_task_modal(self, t_data: dict):
        from gui.task_dialog import TaskDialog
        dialog = TaskDialog(self, title="Edit Task", task_data=t_data)
        result = dialog.get_result()
        if result:
            try:
                conn = self.db.get_connection()
                conn.execute(
                    "UPDATE tasks SET title = ?, task_type = ?, description = ?, updated_at = ? WHERE id = ?",
                    (result["title"], result["task_type"], result["description"], 
                     datetime.now(timezone.utc).isoformat(), t_data["id"])
                )
                conn.commit()
                self._refresh_client_sidebar()
                self._load_table("tasks")
                self._set_status("Task updated successfully.")
            except Exception as e:
                messagebox.showerror("Error", str(e))

    def _on_rerun_selected(self):
        tid = self._selected_task_id
        if self._current_table == "tasks":
            sel = self._tree.selection()
            if sel: tid = int(sel[0])
        
        if not tid:
            messagebox.showinfo("Info", "Select a task first (sidebar or task list).")
            return
        self._on_rerun_task(tid)

    def _on_new_client(self):
        prev = self._current_table
        self._current_table = "clients"
        try:
            self._show_edit_form(row_data=None)
        finally:
            self._current_table = prev

    def _run_task(self, task_id: int):
        messagebox.showinfo("Run", f"Running task {task_id} is not fully implemented yet.")

    def _create_snapshot(self, task_id: int):
        messagebox.showinfo("Info", "Manual snapshot creation from staging is disabled. Use the scraper to create new snapshots.")


    # ──────────────── data loading ────────────────────────────────────────────────────

    def _load_table(self, key: str):
        if key not in TABLES:
            return
        self._current_table = key
        td = TABLES[key]
        cols = td["cols"]
        
        if hasattr(self, "_details_panel"):
            self._details_panel.hide()

        # Enable Export Report only on Snapshots view
        try:
            new_state = "normal" if key == "snapshots" else "disabled"
            self._export_report_btn.configure(state=new_state)
        except Exception:
            pass

        # Reconfigure treeview columns
        visible = [c for c in cols if c[2] > 0]
        col_ids = [c[0] for c in visible]
        self._col_keys = col_ids
        self._tree.configure(columns=col_ids)

        for key_c, header, width, _ in visible:
            self._tree.heading(
                key_c, text=header,
                command=lambda c=key_c: self._sort_column(c))
            self._tree.column(key_c, width=width, minwidth=30, stretch=(width > 150))

        self._rebuild_col_filters(cols)
        self._global_var.set("")
        self._sort_col = None
        self._sort_dir = 1
        self._fetch_and_display()

    def _fetch_and_display(self):
        td = TABLES[self._current_table]
        sql = td["sql"]
        params: list = []

        # Inject foreign key filter if applicable
        fk_join = td.get("fk_join")
        filter_val = None
        if fk_join:
            if "client_id" in fk_join:
                filter_val = self._selected_client_id
            elif "task_id" in fk_join:
                filter_val = self._selected_task_id
            elif "snapshot_id" in fk_join:
                filter_val = self._selected_snapshot_id

        if fk_join and filter_val is not None:
            low = sql.lower()
            
            # Find the best place to insert (before ORDER BY or LIMIT)
            insert_pos = len(sql)
            for marker in [" order by ", " limit "]:
                pos = low.find(marker)
                if pos != -1 and pos < insert_pos:
                    insert_pos = pos
            
            main_sql = sql[:insert_pos]
            suffix = sql[insert_pos:]
            
            # Use a more robust check for top-level WHERE (not inside subqueries)
            has_where = False
            paren_level = 0
            low_main = main_sql.lower()
            for i in range(len(low_main)):
                if low_main[i] == '(':
                    paren_level += 1
                elif low_main[i] == ')':
                    paren_level -= 1
                elif paren_level == 0 and low_main[i:i+7] == " where ":
                    has_where = True
                    break
            
            if has_where:
                main_sql += f" AND {fk_join} = ?"
            else:
                main_sql += f" WHERE {fk_join} = ?"
            
            sql = main_sql + suffix
            params.append(filter_val)

        try:
            conn = self.db.get_connection()
            rows = conn.execute(sql, params).fetchall()
            self._all_rows = [dict(r) for r in rows]
        except Exception as exc:
            logger.error("[DBPanel] Query error: %s", exc)
            self._all_rows = []
            self._set_status(f"Query error: {exc}")

        self._apply_filters()

    def _apply_filters(self, *_):
        global_q = self._global_var.get().lower().strip()
        col_filters = {k: v.get().lower().strip() for k, v in self._filter_vars.items() if v.get().strip()}

        result = []
        for row in self._all_rows:
            # Global search
            if global_q:
                combined = " ".join(str(v) for v in row.values() if v is not None).lower()
                if global_q not in combined:
                    continue
            # Per-column
            skip = False
            for col, fval in col_filters.items():
                cell = str(row.get(col, "") or "").lower()
                if fval not in cell:
                    skip = True
                    break
            if skip:
                continue
            result.append(row)

        # Sort
        if self._sort_col:
            def sort_key(r):
                val = r.get(self._sort_col)
                if val is None: return (True, 0)
                try: 
                    # Attempt numeric cast for proper math sorting
                    return (False, float(val))
                except (ValueError, TypeError):
                    return (False, str(val).lower())

            try:
                result.sort(key=sort_key, reverse=(self._sort_dir == -1))
            except Exception:
                pass

        self._displayed_rows = result
        self._populate_treeview()

    def _populate_treeview(self):
        self._tree.delete(*self._tree.get_children())
        td = TABLES[self._current_table]
        visible_keys = self._col_keys

        for i, row in enumerate(self._displayed_rows):
            vals = [row.get(k, "") for k in visible_keys]
            tag = ""
            # Highlight availability
            if "avail_code" in row and row["avail_code"] == 0:
                tag = "irrelevant"
            elif row.get("rel") == "✓" or row.get("is_relevant") == 0:
                tag = "irrelevant"
            iid = str(row.get(td["pk"], i))
            self._tree.insert("", "end", iid=iid, values=vals, tags=(tag,))

        total = len(self._all_rows)
        shown = len(self._displayed_rows)
        table_label = TABLES[self._current_table]["label"]
        self._stats_lbl.configure(
            text=f"{table_label}  |  {shown} / {total} rows")
        self._set_status(f"Loaded {shown} rows" + (f" (filtered from {total})" if shown != total else ""))

    def _on_filter_change(self, *_):
        self.after(200, self._apply_filters)

    # ──────────────── sorting ─────────────────────────────────────────────────────────

    def _sort_column(self, col: str):
        if self._sort_col == col:
            self._sort_dir = -1 if self._sort_dir == 1 else (0 if self._sort_dir == -1 else 1)
            if self._sort_dir == 0:
                self._sort_col = None
        else:
            self._sort_col = col
            self._sort_dir = 1

        # Update header text
        td = TABLES[self._current_table]
        for key_c, header, _, _ in td["cols"]:
            if key_c not in self._col_keys:
                continue
            indicator = ""
            if self._sort_col == key_c:
                indicator = " ▲" if self._sort_dir == 1 else " ▼"
            try:
                self._tree.heading(key_c, text=header + indicator)
            except Exception:
                pass

        self._apply_filters()

    # ──────────────── toolbar actions ─────────────────────────────────────────────────

    def _on_add_row(self):
        td = TABLES[self._current_table]
        if not td.get("editable"):
            messagebox.showinfo("Read-only", "This table is read-only. Rows are created by the scraper.")
            return
        self._show_edit_form(row_data=None)

    def _on_edit_row(self, event=None):
        sel = self._tree.selection()
        if not sel:
            return
        td = TABLES[self._current_table]
        pk = td["pk"]
        row_id = sel[0]
        row_data = next((r for r in self._displayed_rows if str(r.get(pk)) == str(row_id)), None)

        if not td.get("editable"):
            messagebox.showinfo("Read-only", "This table is read-only. Data is managed automatically by the scraper.")
            return

        if row_data:
            self._show_edit_form(row_data=row_data)

    def _on_double_click(self, event=None):
        """Navigate on double click for tasks/snapshots, else edit."""
        sel = self._tree.selection()
        if not sel: return
        row_id = int(sel[0])

        if self._current_table == "tasks":
            self._on_nav_task(row_id)
        elif self._current_table == "snapshots":
            self._on_nav_snapshot(row_id)
        else:
            self._on_edit_row(event)

    def _on_tree_select(self, event):
        if self._current_table not in ("all_products", "snapshot_products"):
            self._details_panel.hide()
            return
        sel = self._tree.selection()
        if not sel:
            self._details_panel.hide()
            return
        row_id = sel[0]
        td = TABLES[self._current_table]
        pk = td["pk"]
        row_data = next((r for r in self._displayed_rows if str(r.get(pk)) == str(row_id)), None)
        if row_data:
            self._details_panel.show(row_data)
        else:
            self._details_panel.hide()

    def _on_delete_selected(self):
        sel = self._tree.selection()
        if not sel:
            messagebox.showwarning("Nothing selected", "Select one or more rows first.")
            return
        td = TABLES[self._current_table]
        n = len(sel)
        if not messagebox.askyesno("Confirm Delete",
                                   f"Delete {n} selected row(s) from '{td['label']}'?\nThis cannot be undone."):
            return
        ids = []
        for iid in sel:
            try:
                ids.append(int(iid))
            except ValueError:
                pass
        try:
            conn = self.db.get_connection()
            placeholders = ",".join(["?"] * len(ids))
            res = conn.execute(f"DELETE FROM {td['table']} WHERE {td['pk']} IN ({placeholders})", ids)
            conn.commit()
            deleted = res.rowcount
            self._set_status(f"Deleted {deleted} row(s).")
        except Exception as exc:
            messagebox.showerror("Error", str(exc))
        self._fetch_and_display()

    def _on_clear_table(self):
        td = TABLES[self._current_table]
        # Check context
        filtering = None
        if self._current_table == "snapshots" and self._selected_task_id:
            filtering = ("task_id", self._selected_task_id)
            msg = f"Delete all snapshots for the ACTIVE TASK (#{self._selected_task_id})?"
        elif self._current_table == "snapshot_products" and self._selected_snapshot_id:
            filtering = ("snapshot_id", self._selected_snapshot_id)
            msg = f"Delete all products from the ACTIVE SNAPSHOT (#{self._selected_snapshot_id})?"
        else:
            msg = f"Delete ALL rows from '{td['label']}' ({td['table']})?\n\nThis cannot be undone!"

        if not messagebox.askyesno("⚠ Clear Table", msg, icon="warning"):
            return

        try:
            conn = self.db.get_connection()
            if filtering:
                res = conn.execute(f"DELETE FROM {td['table']} WHERE {filtering[0]} = ?", (filtering[1],))
                n = res.rowcount
            else:
                conn = self.db.get_connection()
                res = conn.execute(f"DELETE FROM {td['table']}")
                conn.commit()
                n = res.rowcount
            conn.commit()
            
            self._set_status(f"Cleared {n} rows.")
            self._fetch_and_display()
            if td["section"] == "project":
                self._refresh_client_sidebar()
        except Exception as e:
            messagebox.showerror("Error", str(e))



    def _on_export_report(self):
        """Export a styled multi-sheet Excel report for 2+ selected snapshots."""
        if self._current_table != "snapshots":
            messagebox.showinfo("Info", "Navigate to the Snapshots view and select 2+ snapshots.")
            return

        sel = self._tree.selection()
        if len(sel) < 2:
            messagebox.showwarning(
                "Export Report",
                "Please select at least 2 snapshots to generate a comparison report."
            )
            return

        try:
            snap_ids = sorted(int(i) for i in sel)  # oldest → newest by ID
        except ValueError:
            messagebox.showerror("Error", "Could not parse selected snapshot IDs.")
            return

        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            title="Save Snapshot Report",
            initialfile=f"snapshot_report_{snap_ids[0]}_to_{snap_ids[-1]}.xlsx",
        )
        if not path:
            return

        try:
            from reports.snapshot_report import generate_snapshot_report
            generate_snapshot_report(snap_ids, path, str(self.db.db_path))
            self._set_status(f"Report saved: {path}")
            messagebox.showinfo(
                "Report Exported",
                f"Report generated successfully:\n{path}"
            )
        except ValueError as ve:
            messagebox.showerror("Validation Error", str(ve))
        except Exception as exc:
            logger.exception("[DBPanel] Report generation failed: %s", exc)
            messagebox.showerror("Export Error", f"Report generation failed:\n{exc}")

    def _on_compare_snapshots(self):
        if self._current_table != "snapshots":
            messagebox.showinfo("Info", "Select 2 records in the Snapshots table to compare.")
            return
            
        sel = self._tree.selection()
        if len(sel) != 2:
            messagebox.showwarning("Compare Snapshots", "Please select exactly 2 snapshots to compare.")
            return
            
        try:
            s_ids = [int(i) for i in sel]
        except ValueError:
            return
        
        s_ids.sort()
        s1, s2 = s_ids[0], s_ids[1]

        DiffPanel(self, self.db, s1, s2).show()

    def _on_normalize(self):
        if self._current_table != "all_products" or not self.scheduler:
            messagebox.showinfo("Info", "Normalization is only available for All Products.")
            return
        self._set_status("AI Normalization running…")

        def work():
            import asyncio
            loop = asyncio.new_event_loop()
            asyncio.set_event_loop(loop)
            try:
                loop.run_until_complete(
                    self.scheduler.normalize_all_pending(stop_event=self.scheduler._stop_event))
            finally:
                loop.close()
            self.after(0, lambda: (self._fetch_and_display(),
                                   self._set_status("Normalization complete.")))

        threading.Thread(target=work, daemon=True).start()

    def _on_edit_prompt(self):
        """Open a window to view and edit the AI normalization prompt."""
        PromptEditorDialog(self).show()

    def _on_export(self):
        if not self._displayed_rows:
            messagebox.showinfo("Empty", "Nothing to export.")
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            title="Export current view")
        if not path:
            return
        try:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = TABLES[self._current_table]["label"][:30]

            # Header
            headers = [c[1] for c in TABLES[self._current_table]["cols"] if c[2] > 0]
            keys = self._col_keys
            ws.append(headers)

            for row in self._displayed_rows:
                ws.append([row.get(k, "") for k in keys])

            # Auto-width
            for col_cells in ws.columns:
                length = max(len(str(c.value or "")) for c in col_cells)
                ws.column_dimensions[col_cells[0].column_letter].width = min(length + 2, 60)

            wb.save(path)
            messagebox.showinfo("Exported", f"Saved {len(self._displayed_rows)} rows to:\n{path}")
        except Exception as exc:
            messagebox.showerror("Export error", str(exc))

    def _on_refresh(self):
        self._refresh_client_sidebar()
        self._fetch_and_display()
        self._set_status("Refreshed.")

    # ──────────────── context menu ────────────────────────────────────────────────────

    def _show_context_menu(self, event):
        item = self._tree.identify_row(event.y)
        if not item:
            return
        if item not in self._tree.selection():
            self._tree.selection_set(item)

        menu = tk.Menu(self, tearoff=0,
                       background="#2d2d2d", foreground="#cccccc",
                       activebackground="#1f538d", activeforeground="#ffffff")

        # URL actions if there's a url column
        row = next((r for r in self._displayed_rows
                    if str(r.get(TABLES[self._current_table]["pk"])) == item), None)
        url = (row or {}).get("url") or (row or {}).get("product_url") or \
              (row or {}).get("seller_url") or (row or {}).get("competitor_product_url") or \
              (row or {}).get("comp_url") or (row or {}).get("output_path")

        if url and url.startswith("http"):
            menu.add_command(label="🌐  Open URL", command=lambda: webbrowser.open(url))
            menu.add_command(label="📋  Copy URL", command=lambda: (
                self.clipboard_clear(), self.clipboard_append(url)))
            menu.add_separator()

        menu.add_command(label="📝  Edit Row", command=self._on_edit_row)
        menu.add_command(label="🗑  Delete Row", command=self._on_delete_selected)
        menu.tk_popup(event.x_root, event.y_root)

    # ──────────────── edit / add form ─────────────────────────────────────────────────

    def _show_edit_form(self, row_data: Optional[dict]):
        td = TABLES[self._current_table]
        EditFormDialog(self, self.db, td, row_data, self._on_edit_saved).show()

    def _on_edit_saved(self, is_project_section: bool):
        if is_project_section:
            self._refresh_client_sidebar()
        self._fetch_and_display()




